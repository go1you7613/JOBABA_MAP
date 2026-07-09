#!/usr/bin/env python3
"""Load docs/data CSV files into the local jobaba_map MariaDB database.

The loader reads the provided logical/physical data element workbooks, creates
or extends only the source tables represented by docs/data/*.csv, and upserts
CSV rows by the documented primary keys.
"""

from __future__ import annotations

import csv
import os
import re
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DATA = DOCS / "data"
LOGICAL_XLSX = next(DOCS.glob("25NIA-DE5057-01_*v1.0.xlsx"))
PHYSICAL_XLSX = next(DOCS.glob("25NIA-DE5059-01_*v1.0.xlsx"))

MYSQL = os.environ.get("MYSQL_BIN", "/opt/homebrew/bin/mysql")
DB_HOST = os.environ.get("JOBABA_DB_HOST", "localhost")
DB_PORT = os.environ.get("JOBABA_DB_PORT", "3306")
DB_USER = os.environ.get("JOBABA_DB_USER", "jobaba")
DB_PASSWORD = os.environ.get("JOBABA_DB_PASSWORD")
DB_NAME = os.environ.get("JOBABA_DB_NAME", "jobaba_map")

if not DB_PASSWORD:
    raise RuntimeError("JOBABA_DB_PASSWORD is required")


@dataclass
class ColumnDef:
    name: str
    sql_type: str
    nullable: bool
    default: Optional[str] = None
    comment: str = ""
    pk: bool = False


@dataclass
class CsvTarget:
    csv_path: Path
    table_name: str
    headers: List[str]
    row_count: int
    encoding: str


def qident(name: str) -> str:
    return "`" + name.replace("`", "``") + "`"


def sql_string(value: str) -> str:
    return "'" + value.replace("\\", "\\\\").replace("'", "''") + "'"


def normalize_table_name(csv_path: Path) -> str:
    name = csv_path.name
    match = re.search(r"\(\s*([^)]+?)\s*\)", name)
    if match:
        raw = match.group(1).strip().replace(" ", "")
    elif re.fullmatch(r"[A-Za-z0-9_.]+", csv_path.stem):
        raw = csv_path.stem
    elif "공통코드" in name:
        raw = "jedut.tb_code"
    else:
        raw = csv_path.stem
    return raw.split(".")[-1].lower()


def read_csv_header_and_count(path: Path) -> Tuple[str, List[str], int]:
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            with path.open("r", encoding=enc, newline="") as f:
                reader = csv.reader(f)
                header = next(reader)
                count = sum(1 for _ in reader)
            return enc, [h.strip() for h in header], count
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("csv", b"", 0, 1, f"unsupported encoding: {path}")


def load_csv_targets() -> List[CsvTarget]:
    targets = []
    for path in sorted(DATA.glob("*.csv")):
        enc, headers, count = read_csv_header_and_count(path)
        targets.append(
            CsvTarget(
                csv_path=path,
                table_name=normalize_table_name(path),
                headers=headers,
                row_count=count,
                encoding=enc,
            )
        )
    return targets


def parse_type(type_name: str, type_len: str) -> str:
    t = (type_name or "").strip().lower()
    length = (type_len or "").strip().lower()
    if length and length not in {"null", "none"}:
        if "(" in length or length in {"text", "mediumtext", "longtext", "datetime", "date", "bigint", "int"}:
            return length
    if t in {"varchar", "char"}:
        size = re.search(r"\d+", length)
        return f"{t}({size.group(0) if size else 255})"
    if t in {"int", "bigint", "smallint", "tinyint"}:
        return t
    if t in {"datetime", "date", "time"}:
        return t
    if t in {"text", "mediumtext", "longtext"}:
        return t
    return "text"


def read_workbook_defs(path: Path) -> Dict[str, Dict[str, ColumnDef]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tables: Dict[str, Dict[str, ColumnDef]] = {}
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            continue
        headers = [str(v).strip() if v is not None else "" for v in header_row]
        if "테이블명" in headers:
            table_idx = headers.index("테이블명")
            column_idx = headers.index("컬럼명")
            type_idx = headers.index("컬럼타입")
            len_idx = headers.index("컬럼길이(LENGTH)")
            null_idx = headers.index("NULL여부")
            pk_idx = headers.index("컬럼키(PK)") if "컬럼키(PK)" in headers else None
            default_idx = headers.index("디폴트값") if "디폴트값" in headers else None
            comment_idx = headers.index("테이블 코멘트") if "테이블 코멘트" in headers else None
        elif "물리테이블" in headers:
            table_idx = headers.index("물리테이블")
            column_idx = headers.index("물리컬럼")
            type_idx = headers.index("물리타입")
            len_idx = headers.index("길이")
            null_idx = headers.index("NULL여부")
            pk_idx = None
            default_idx = None
            comment_idx = headers.index("컬럼설명") if "컬럼설명" in headers else None
        else:
            continue

        for row in rows:
            table = str(row[table_idx]).strip().lower() if row[table_idx] is not None else ""
            column = str(row[column_idx]).strip() if row[column_idx] is not None else ""
            if not table or not column:
                continue
            pk_value = str(row[pk_idx]).strip().upper() if pk_idx is not None and row[pk_idx] is not None else ""
            default = str(row[default_idx]).strip() if default_idx is not None and row[default_idx] is not None else None
            comment = str(row[comment_idx]).strip() if comment_idx is not None and row[comment_idx] is not None else ""
            tables.setdefault(table, {})[column.upper()] = ColumnDef(
                name=column.upper(),
                sql_type=parse_type(str(row[type_idx] or ""), str(row[len_idx] or "")),
                nullable=str(row[null_idx]).strip().upper() != "NO",
                default=default,
                comment=comment,
                pk=pk_value == "PRI",
            )
    return tables


def col_sql(col: ColumnDef, *, for_create: bool) -> str:
    parts = [qident(col.name), col.sql_type]
    parts.append("NULL" if col.nullable else "NOT NULL")
    if for_create and col.default:
        default_upper = col.default.upper()
        if default_upper in {"CURRENT_TIMESTAMP", "NULL"}:
            parts.append(f"DEFAULT {default_upper}")
        elif col.default != "":
            parts.append("DEFAULT " + sql_string(col.default))
    if col.comment:
        parts.append("COMMENT " + sql_string(col.comment[:1024]))
    return " ".join(parts)


def build_table_defs(targets: Iterable[CsvTarget]) -> Dict[str, List[ColumnDef]]:
    logical = read_workbook_defs(LOGICAL_XLSX)
    physical = read_workbook_defs(PHYSICAL_XLSX)
    table_defs: Dict[str, List[ColumnDef]] = {}
    for target in targets:
        source = logical.get(target.table_name) or physical.get(target.table_name)
        if source is None:
            raise RuntimeError(f"No workbook definition for {target.table_name}")
        phys = physical.get(target.table_name, {})
        cols = []
        missing = []
        for header in target.headers:
            key = header.upper()
            col = source.get(key) or phys.get(key)
            if col is None:
                missing.append(header)
                continue
            merged = ColumnDef(
                name=key,
                sql_type=col.sql_type,
                nullable=col.nullable,
                default=col.default,
                comment=col.comment,
                pk=phys.get(key, col).pk,
            )
            cols.append(merged)
        if missing:
            raise RuntimeError(f"{target.table_name}: missing definitions for {missing}")
        table_defs[target.table_name] = cols
    return table_defs


def create_schema_sql(table: str, cols: List[ColumnDef]) -> str:
    pk_cols = [col.name for col in cols if col.pk]
    lines = [f"CREATE TABLE IF NOT EXISTS {qident(table)} ("]
    col_lines = ["  " + col_sql(col, for_create=True) for col in cols]
    if pk_cols:
        col_lines.append("  PRIMARY KEY (" + ", ".join(qident(c) for c in pk_cols) + ")")
    lines.append(",\n".join(col_lines))
    lines.append(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;")
    for col in cols:
        lines.append(
            f"ALTER TABLE {qident(table)} ADD COLUMN IF NOT EXISTS {col_sql(col, for_create=False)};"
        )
    return "\n".join(lines)


def mysql_base_cmd(*extra: str) -> List[str]:
    return [
        MYSQL,
        "-h",
        DB_HOST,
        "-P",
        DB_PORT,
        "-u",
        DB_USER,
        f"--password={DB_PASSWORD}",
        "--default-character-set=utf8mb4",
        "--local-infile=1",
        DB_NAME,
        *extra,
    ]


def run_mysql(sql: str) -> str:
    proc = subprocess.run(
        mysql_base_cmd("-N", "-B", "-e", sql),
        check=True,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    return proc.stdout


def execute_sql_file(path: Path) -> None:
    with path.open("r", encoding="utf-8") as f:
        subprocess.run(mysql_base_cmd(), check=True, text=True, stdin=f)


def normalize_value(value: str, col: ColumnDef) -> str:
    value = (value or "").strip()
    if value == "":
        return r"\N" if col.nullable else ""
    typ = col.sql_type.lower()
    if "datetime" in typ:
        match = re.fullmatch(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})(?:\s+(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?)?", value)
        if match:
            y, mo, d, h, mi, s = match.groups()
            return f"{int(y):04d}-{int(mo):02d}-{int(d):02d} {int(h or 0):02d}:{int(mi or 0):02d}:{int(s or 0):02d}"
    if typ in {"int", "bigint", "smallint", "tinyint"}:
        if re.fullmatch(r"-?\d+", value):
            return value
        return r"\N" if col.nullable else "0"
    return value


def write_normalized_tsv(target: CsvTarget, cols: List[ColumnDef], output_dir: Path) -> Path:
    col_by_name = {col.name: col for col in cols}
    out = output_dir / f"{target.table_name}.tsv"
    with target.csv_path.open("r", encoding=target.encoding, newline="") as src, out.open(
        "w", encoding="utf-8", newline=""
    ) as dst:
        reader = csv.DictReader(src)
        writer = csv.writer(dst, delimiter="\t", lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
        for row in reader:
            writer.writerow([normalize_value(row.get(h, ""), col_by_name[h.upper()]) for h in target.headers])
    return out


def build_load_sql(target: CsvTarget, cols: List[ColumnDef], tsv_path: Path) -> str:
    assignments = ", ".join(f"{qident(col.name)} = VALUES({qident(col.name)})" for col in cols)
    return f"""
LOAD DATA LOCAL INFILE {sql_string(str(tsv_path))}
REPLACE INTO TABLE {qident(target.table_name)}
CHARACTER SET utf8mb4
FIELDS TERMINATED BY '\\t'
OPTIONALLY ENCLOSED BY '"'
ESCAPED BY '\\\\'
LINES TERMINATED BY '\\n'
({", ".join(qident(col.name) for col in cols)});
"""


def existing_counts(targets: Iterable[CsvTarget]) -> Dict[str, Optional[int]]:
    counts: Dict[str, Optional[int]] = {}
    for target in targets:
        exists = run_mysql(
            "SELECT COUNT(*) FROM information_schema.tables "
            f"WHERE table_schema = DATABASE() AND table_name = {sql_string(target.table_name)};"
        ).strip()
        if exists == "1":
            cnt = run_mysql(f"SELECT COUNT(*) FROM {qident(target.table_name)};").strip()
            counts[target.table_name] = int(cnt)
        else:
            counts[target.table_name] = None
    return counts


def main() -> int:
    targets = load_csv_targets()
    table_defs = build_table_defs(targets)
    before = existing_counts(targets)

    with tempfile.TemporaryDirectory(prefix="jobaba_docs_load_") as tmp:
        tmp_path = Path(tmp)
        ddl_path = tmp_path / "load_docs_data.sql"
        statements = ["SET SESSION sql_mode = '';", "SET FOREIGN_KEY_CHECKS = 0;"]
        for target in targets:
            cols = table_defs[target.table_name]
            statements.append(create_schema_sql(target.table_name, cols))
            tsv = write_normalized_tsv(target, cols, tmp_path)
            statements.append(build_load_sql(target, cols, tsv))
        statements.append("SET FOREIGN_KEY_CHECKS = 1;")
        ddl_path.write_text("\n\n".join(statements), encoding="utf-8")
        execute_sql_file(ddl_path)

    after = existing_counts(targets)
    print("table,csv_rows,before_db_rows,after_db_rows,loaded_or_updated")
    for target in targets:
        before_count = before[target.table_name]
        after_count = after[target.table_name]
        loaded = after_count - (before_count or 0)
        print(
            f"{target.table_name},{target.row_count},"
            f"{'' if before_count is None else before_count},{after_count},{loaded}"
        )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except subprocess.CalledProcessError as exc:
        print(exc.stderr or str(exc), file=sys.stderr)
        raise
